"""CSV and Excel artefact register export (plan task 7.4).

Produces a flat table — one row per artefact — suitable for GRC teams who
live in spreadsheets.  The CSV exporter returns ``str``; the Excel exporter
returns ``bytes`` (XLSX).

Columns cover every reportable field from the artefact, its risk assessment
and its PQC recommendation so the register is self-contained.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from app.schemas.artefact import (
    AlgorithmDetail,
    CertificateDetail,
    CryptoArtefact,
    KeyDetail,
)
from app.schemas.recommendation import PqcRecommendation
from app.schemas.risk import RiskAssessment
from app.schemas.scan import ScanResult

#: Column order for the flat register.
COLUMNS: list[str] = [
    "artefact_id",
    "name",
    "asset_type",
    "algorithm",
    "primitive",
    "key_size_bits",
    "mode",
    "padding",
    "curve",
    "quantum_vulnerability",
    "nist_security_level",
    "location",
    "discovered_by",
    "first_seen",
    "last_seen",
    "risk_score",
    "priority",
    "assessment_status",
    "mosca_urgent",
    "mosca_shortfall_years",
    "mosca_deadline_year",
    "recommended_algorithm",
    "recommended_parameter_set",
    "is_hybrid",
    "migration_wave",
]


def _flatten_artefact(
    artefact: CryptoArtefact,
    assessment: RiskAssessment | None,
    recommendation: PqcRecommendation | None,
) -> dict[str, Any]:
    """Flatten an artefact + assessment + recommendation into a row dict."""
    detail = artefact.detail

    key_size = getattr(detail, "key_size_bits", None) or getattr(
        detail, "size_bits", None
    )
    if key_size is None and isinstance(detail, CertificateDetail):
        key_size = detail.public_key_size_bits

    mode = (
        detail.mode.value
        if isinstance(detail, AlgorithmDetail) and detail.mode
        else None
    )
    padding = (
        detail.padding.value
        if isinstance(detail, AlgorithmDetail) and detail.padding
        else None
    )
    curve = getattr(detail, "curve", None)
    if curve is None and isinstance(detail, CertificateDetail):
        curve = detail.public_key_curve

    row: dict[str, Any] = {
        "artefact_id": artefact.artefact_id,
        "name": artefact.name,
        "asset_type": artefact.asset_type.value,
        "algorithm": artefact.algorithm or "",
        "primitive": artefact.primitive.value if artefact.primitive else "",
        "key_size_bits": key_size or "",
        "mode": mode or "",
        "padding": padding or "",
        "curve": curve or "",
        "quantum_vulnerability": artefact.quantum_vulnerability.value,
        "nist_security_level": artefact.nist_security_level.value,
        "location": artefact.primary_location,
        "discovered_by": artefact.discovered_by.value,
        "first_seen": artefact.first_seen.isoformat(),
        "last_seen": artefact.last_seen.isoformat(),
        "risk_score": "",
        "priority": "",
        "assessment_status": "",
        "mosca_urgent": "",
        "mosca_shortfall_years": "",
        "mosca_deadline_year": "",
        "recommended_algorithm": "",
        "recommended_parameter_set": "",
        "is_hybrid": "",
        "migration_wave": "",
    }

    if assessment:
        row["risk_score"] = (
            assessment.final_score if assessment.final_score is not None else ""
        )
        row["priority"] = assessment.priority.value
        row["assessment_status"] = assessment.status.value
        row["mosca_urgent"] = (
            str(assessment.mosca.is_urgent).lower()
            if assessment.mosca.is_urgent is not None
            else ""
        )
        row["mosca_shortfall_years"] = (
            assessment.mosca.shortfall_years
            if assessment.mosca.shortfall_years is not None
            else ""
        )
        row["mosca_deadline_year"] = (
            assessment.resource.migration_deadline_year
            if assessment.resource.migration_deadline_year is not None
            else ""
        )

    if recommendation:
        row["recommended_algorithm"] = (
            recommendation.recommended_algorithm or ""
        )
        row["recommended_parameter_set"] = (
            recommendation.recommended_parameter_set or ""
        )
        row["is_hybrid"] = str(recommendation.is_hybrid).lower()

    return row


def _build_rows(
    scan_result: ScanResult,
    assessments: list[RiskAssessment],
    recommendations: list[PqcRecommendation],
) -> list[dict[str, Any]]:
    """Build the flat row list from scan data."""
    assessment_by_artefact = {a.artefact_id: a for a in assessments}
    recommendation_by_assessment = {r.assessment_id: r for r in recommendations}

    rows: list[dict[str, Any]] = []
    for artefact in scan_result.deduplicated():
        assessment = assessment_by_artefact.get(artefact.artefact_id)
        rec = (
            recommendation_by_assessment.get(assessment.assessment_id)
            if assessment
            else None
        )
        rows.append(_flatten_artefact(artefact, assessment, rec))
    return rows


def export_csv(
    scan_result: ScanResult,
    assessments: list[RiskAssessment] | None = None,
    recommendations: list[PqcRecommendation] | None = None,
) -> str:
    """Export the artefact register as CSV."""
    assessments = assessments or []
    recommendations = recommendations or []
    rows = _build_rows(scan_result, assessments, recommendations)

    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=COLUMNS, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return buf.getvalue()


def export_excel(
    scan_result: ScanResult,
    assessments: list[RiskAssessment] | None = None,
    recommendations: list[PqcRecommendation] | None = None,
) -> bytes:
    """Export the artefact register as an XLSX workbook.

    Requires ``openpyxl`` (optional ``report`` dependency group).
    Features: conditional formatting by priority, auto-filter, frozen header
    row, auto-sized columns, and a summary sheet.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ImportError(
            "Excel export requires 'openpyxl'. "
            "Install with: pip install trinetra-backend[report]"
        ) from exc

    assessments = assessments or []
    recommendations = recommendations or []
    rows = _build_rows(scan_result, assessments, recommendations)

    wb = Workbook()

    # -- Inventory sheet ---------------------------------------------------
    ws = wb.active
    ws.title = "Artefact Inventory"  # type: ignore[union-attr]

    # Header
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2C3E50")
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)  # type: ignore[union-attr]
        cell.font = header_font
        cell.fill = header_fill

    # Data
    priority_fills = {
        "p0": PatternFill("solid", fgColor="FF6B6B"),
        "p1": PatternFill("solid", fgColor="FECA57"),
        "p2": PatternFill("solid", fgColor="48DBFB"),
    }

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, col_name in enumerate(COLUMNS, 1):
            cell = ws.cell(  # type: ignore[union-attr]
                row=row_idx, column=col_idx, value=row_data.get(col_name, "")
            )
            # Conditional fill for priority column
            if col_name == "priority" and row_data.get("priority") in priority_fills:
                cell.fill = priority_fills[row_data["priority"]]

    # Auto-filter and freeze
    ws.auto_filter.ref = (  # type: ignore[union-attr]
        f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
    )
    ws.freeze_panes = "A2"  # type: ignore[union-attr]

    # Auto-size columns (approximate)
    for col_idx, col_name in enumerate(COLUMNS, 1):
        max_len = len(col_name)
        for row_data in rows[:50]:  # Sample first 50 rows
            val = str(row_data.get(col_name, ""))
            max_len = max(max_len, min(len(val), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3  # type: ignore[union-attr]

    # -- Summary sheet -----------------------------------------------------
    summary = wb.create_sheet("Summary")
    summary_font = Font(bold=True)

    summary.cell(row=1, column=1, value="Trinetra Artefact Summary").font = (
        Font(bold=True, size=14)
    )
    summary.cell(
        row=2,
        column=1,
        value=f"Scan: {scan_result.scan_id}",
    )
    summary.cell(
        row=3,
        column=1,
        value=f"Target: {scan_result.target.identifier}",
    )
    summary.cell(
        row=4,
        column=1,
        value=f"Total artefacts: {len(rows)}",
    )

    # Counts by type
    row_offset = 6
    summary.cell(row=row_offset, column=1, value="By Asset Type").font = summary_font
    type_counts: dict[str, int] = {}
    for row_data in rows:
        t = str(row_data.get("asset_type", "unknown"))
        type_counts[t] = type_counts.get(t, 0) + 1
    for i, (asset_type, count) in enumerate(sorted(type_counts.items()), 1):
        summary.cell(row=row_offset + i, column=1, value=asset_type)
        summary.cell(row=row_offset + i, column=2, value=count)

    # Counts by priority
    row_offset = row_offset + len(type_counts) + 2
    summary.cell(row=row_offset, column=1, value="By Priority").font = summary_font
    priority_counts: dict[str, int] = {}
    for row_data in rows:
        p = str(row_data.get("priority", "")) or "unassessed"
        priority_counts[p] = priority_counts.get(p, 0) + 1
    for i, (priority, count) in enumerate(sorted(priority_counts.items()), 1):
        summary.cell(row=row_offset + i, column=1, value=priority)
        summary.cell(row=row_offset + i, column=2, value=count)

    # Write to bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["export_csv", "export_excel"]
